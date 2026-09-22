from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

out=Path(__file__).resolve().parents[1]/'output'
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
ink='142C3B';ground='F5F8FA';primary='126F75';muted='506678';accent='DDEEEA'
def box(slide,txt,x,y,w,h,size=24,bold=False,color=ink,fill=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)) if fill else slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    if fill: sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=Inches(.12);tf.margin_right=Inches(.08)
    r=tf.paragraphs[0].add_run();r.text=txt;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
def slide(title,steps,caption,notes,num):
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(ground)
    box(s,title,.7,.65,11.9,1.4,34,True)
    for i,label in enumerate(steps):
        x=.8+i*3.1;box(s,label,x,3.0,2.45,1.2,23,True,fill=accent)
        if i<len(steps)-1:box(s,'→',x+2.5,3.25,.5,.65,26)
    box(s,caption,.8,5.25,11.8,1.15,25,color=primary)
    if num:box(s,str(num),11.8,6.9,.55,.3,12)
    s.notes_slide.notes_text_frame.text=notes+' Source: Voice Math source files read in this task.'
slide('Your app handles math locally and voice through ElevenLabs',['Browser app','Local math','ElevenLabs','Conversation'],'Voice Math / architecture explainer','Overview of the components, not a literal data flow: local math and cloud voice are separate paths. Audience: app owner. Approximate length: 5 minutes.',0)
slide('Selected questions reach the local calculator',['Chrome extension','Browser app','Math endpoint','Answer'],'No credentials are needed for demo math.','Extension adds selected text as a URL query. App posts to /api/calculate. Express evaluates restricted arithmetic using mathjs. Example: 15% of 200 = 30.',2)
slide('Your server authorizes a direct browser voice connection',['Browser','Local server','Session API','Signed URL'],'Browser ↔ ElevenLabs: audio and transcript','The signed URL returns via the local server to the browser. The API key never goes to the browser. Browser then starts a WebSocket conversation directly with the agent.',3)
slide('The voice agent requests accurate calculations through your app',['Agent tool call','Browser callback','Local math','Result to agent'],'The agent speaks after receiving the tool result.','The browser callback posts the expression to /api/calculate. Result returns through browser SDK to the agent. Agent is instructed to use the calculator, not guess.',4)
slide('One setup command creates the agent ID for you',['API key','agent:setup','Create agent','Save agent ID'],'Add your key, run setup, restart the app.','Key goes in .env. npm run agent:setup creates the agent via ElevenLabs API and writes .agent-id. Both API key and agent ID are required internally. No placeholder identity or statistics are invented.',5)
p.save(out/'presentation-generator-voice-math.pptx')
print(out/'presentation-generator-voice-math.pptx')
